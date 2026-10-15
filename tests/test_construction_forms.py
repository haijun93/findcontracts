from io import BytesIO
from zipfile import ZipFile
import json
import unittest
from docx import Document
from construction.core import generate_package, validate_facts, DocumentError
from test_construction import analysis


def full_text(data):
    d=Document(BytesIO(data))
    return '\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)


class SupplementalFormTests(unittest.TestCase):
    def test_all_five_documents_share_reviewed_values(self):
        facts=analysis()['facts']
        facts.update(supplier='시험공사업체', representative='테스트대표', business_number='123-45-67890',
                     authority='테스트구', price_unit='식',price_quantity='1',price_basis='입력 견적과 비교견적 검토 예정',
                     contract_reason='현장 여건 검토 후 작성한 사유',legal_basis='담당자가 확인한 근거',
                     count_year='2025',department_count='0',district_count='3',count_confirmation='확인완료',
                     competitive_product='아니오',enterprise_type='확인된 유형',
                     form_date='2026. 10. 15.',preparer='조사자 직급 성명',verifier='확인자 직위 성명',
                     department_signer='사업부서장 테스트1',finance_signer='계약부서장 테스트2',
                     comparisons=[{'vendor':'비교업체','amount':'3,100,000','note':'동일 범위·부가세 포함','source':'비교견적 1쪽'}])
        with ZipFile(generate_package(facts)) as z:
            names=z.namelist()
            self.assertEqual(len(names),11)
            self.assertEqual(len([n for n in names if n.endswith('.docx')]),5)
            for n in names:
                if n.endswith('.docx'):
                    text=full_text(z.read(n))
                    self.assertIn(facts['title'],text)
                    self.assertIn('초안',text)
            survey=full_text(z.read('산출기초조사서_초안.docx'))
            self.assertIn('3,100,000원',survey)
            self.assertIn('2,970,000원',survey)
            self.assertIn('비교업체',survey)
            reason=full_text(z.read('수의계약사유서_초안.docx'))
            self.assertIn('부서: 0회 / 구 전체: 3회',reason)
            self.assertIn('확인 연도: 2025',reason)
            self.assertIn('현장 여건 검토 후 작성한 사유',reason)
            self.assertIn('123-45-67890',reason)
            pledge=full_text(z.read('청렴계약_이행서약서_초안.docx'))
            self.assertIn('사업부서장 테스트1',pledge)
            self.assertIn('계약부서장 테스트2',pledge)
            self.assertIn('서명 또는 날인',pledge)
            self.assertIn('공무원용',pledge)
            manifest=json.loads(z.read('추출근거_검토기록.json'))
            self.assertEqual(manifest['facts']['supplier'],'시험공사업체')

    def test_example_values_never_copied_or_confirmed(self):
        with ZipFile(generate_package(analysis()['facts'])) as z:
            for name in ['산출기초조사서_초안.docx','수의계약사유서_초안.docx','청렴계약_이행서약서_초안.docx']:
                text=full_text(z.read(name))
                for example in ['이피코리아','에이앤스틸','배문찬','김나래','주하나','문철우','9,360,000','9,000,000','제25조제1항제5호가목','√']:
                    self.assertNotIn(example,text)
                self.assertIn('[확인 필요]',text)
            reason=full_text(z.read('수의계약사유서_초안.docx'))
            self.assertNotIn('부서: 0회',reason)
            self.assertIn('확인이 완료되지 않았습니다',reason)

    def test_supplier_candidates_come_from_quote_not_examples(self):
        from construction.core import extract_document, analyze
        from test_construction import QUOTE, REQUEST
        quote = QUOTE + '\n상호   새공사업체   대표   새대표(인)      상호   발주기관\n등록번호 123-45-67890\n'
        data = analyze(extract_document(quote.encode(),'q.txt','견적서'),extract_document(REQUEST.encode(),'r.txt','의뢰서'))
        self.assertEqual(data['facts']['supplier'],'새공사업체')
        self.assertEqual(data['facts']['representative'],'새대표')
        self.assertEqual(data['facts']['business_number'],'123-45-67890')
        self.assertTrue(any(e['field']=='supplier' for e in data['evidence']))
        self.assertEqual(data['facts']['count_confirmation'],'')

    def test_invalid_form_values_rejected(self):
        for key,value in [('department_count','-1'),('district_count','1.5'),('count_year','2026년'),
                          ('price_quantity','0'),('competitive_product','자동승인'),('count_confirmation','√')]:
            facts=analysis()['facts'];facts[key]=value
            with self.assertRaises(DocumentError):validate_facts(facts)
        facts=analysis()['facts'];facts['comparisons']=[{'vendor':'A','amount':-1}]
        with self.assertRaises(DocumentError):validate_facts(facts)

    def test_inconsistent_contract_counts_flagged(self):
        facts=analysis()['facts'];facts.update(department_count='4',district_count='1')
        with ZipFile(generate_package(facts)) as z:
            self.assertIn('부서 계약횟수가 구 전체 계약횟수보다 큽니다',full_text(z.read('수의계약사유서_초안.docx')))
