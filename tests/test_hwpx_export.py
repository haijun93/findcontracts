from io import BytesIO
import re
import unittest
from zipfile import ZipFile, ZIP_STORED
from lxml import etree
from docx import Document
from hwpx import HwpxDocument, validate_editor_open_safety
from construction.hwpx_export import docx_to_hwpx
from construction.core import generate_package
from test_construction import analysis


def norm(text): return re.sub(r'\s+', '', text)


def hwpx_text(data):
    with ZipFile(BytesIO(data)) as archive:
        root=etree.fromstring(archive.read('Contents/section0.xml'))
        return '\n'.join(root.xpath('//*[local-name()="t"]/text()'))


class HwpxExportTests(unittest.TestCase):
    def test_korean_paragraphs_tables_and_manual_breaks_preserved(self):
        doc=Document();doc.add_heading('시방서\n검토용 초안',0)
        doc.add_paragraph('공사내용: 지주식 안내판 7개 철거')
        table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='위치';table.cell(0,1).text='금액'
        table.cell(1,0).text='어울마당로 155 앞\n전기 마감'
        table.cell(1,1).text='2,970,000원'
        source=BytesIO();doc.save(source)
        data=docx_to_hwpx(source.getvalue())
        with ZipFile(BytesIO(data)) as archive:
            self.assertEqual(archive.read('mimetype'),b'application/hwp+zip')
            self.assertEqual(archive.infolist()[0].filename,'mimetype')
            self.assertEqual(archive.infolist()[0].compress_type,ZIP_STORED)
            root=etree.fromstring(archive.read('Contents/section0.xml'))
            self.assertEqual(len(root.xpath('//*[local-name()="tbl"]')),1)
            self.assertEqual(len(root.xpath('//*[local-name()="tc"]')),4)
        text=norm(hwpx_text(data))
        for phrase in ['시방서','검토용 초안','지주식 안내판 7개 철거','어울마당로 155 앞','전기 마감','2,970,000원']:
            self.assertIn(norm(phrase),text)
        report=validate_editor_open_safety(data)
        self.assertTrue(report.reopen_ok)
        self.assertFalse(report.blocking_package_errors)
        self.assertFalse(report.validate_document.issues)
        self.assertIsNotNone(HwpxDocument.open(data))

    def test_every_generated_document_has_hwpx_equivalent(self):
        facts=analysis()['facts']
        with ZipFile(generate_package(facts)) as archive:
            names=archive.namelist()
            self.assertEqual(len([n for n in names if n.endswith('.hwpx')]),5)
            for name in [n for n in names if n.endswith('.docx')]:
                document=Document(BytesIO(archive.read(name)))
                target=norm(hwpx_text(archive.read(name.replace('.docx','.hwpx'))))
                for p in document.paragraphs:
                    if p.text.strip():self.assertIn(norm(p.text),target)
                for t in document.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            if cell.text.strip():self.assertIn(norm(cell.text),target)
