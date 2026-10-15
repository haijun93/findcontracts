"""Convert text-and-table DOCX drafts to native HWPX packages, without Hancom COM."""
from io import BytesIO
import math
import unicodedata
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from hwpx import HwpxDocument, validate_editor_open_safety


def _lines(text, width, size):
    # Conservative height estimate for wrapped East Asian text; Hancom lays out the actual text.
    capacity = max(1, width / (size * 100))
    return sum(max(1, math.ceil(sum(1 if unicodedata.east_asian_width(c) in 'WF' else .55 for c in line) / capacity))
               for line in (text.splitlines() or ['']))


def docx_to_hwpx(data):
    source = Document(BytesIO(data))
    target = HwpxDocument.new()
    target.package.clear_document_metadata()
    target.package.set_document_metadata(title=source.core_properties.title or '검토용 초안', creator='')
    body_style = target.styles.ensure_run(font='맑은 고딕', size=10)
    title_style = target.styles.ensure_run(font='맑은 고딕', size=18, bold=True)
    heading_style = target.styles.ensure_run(font='맑은 고딕', size=13, bold=True)
    subheading_style = target.styles.ensure_run(font='맑은 고딕', size=11, bold=True)
    table_style = target.styles.ensure_run(font='맑은 고딕', size=9)
    first = True
    preview = []
    section = source.sections[0]
    page_width = int(section.page_width.pt * 100)
    width = page_width - int((section.left_margin.pt + section.right_margin.pt) * 100)
    hp = '{http://www.hancom.co.kr/hwpml/2011/paragraph}'
    page_pr = target.sections[0].element.find('.//' + hp + 'pagePr')
    if page_pr is not None:
        page_pr.set('width', str(page_width)); page_pr.set('height', str(int(section.page_height.pt * 100)))
        margin = page_pr.find(hp + 'margin')
        if margin is not None:
            for key in ('left', 'right', 'top', 'bottom'):
                margin.set(key, str(int(getattr(section, key + '_margin').pt * 100)))

    def paragraph(text, style, page_break=False):
        nonlocal first
        # Each DOCX manual line becomes an actual HWPX paragraph; no text is dropped.
        for index, line in enumerate(text.split('\n')):
            if first:
                para = target.paragraphs[0]; para.text = line
                for run in para.runs: run.element.set('charPrIDRef', style)
                first = False
            else:
                para = target.add_paragraph(line, char_pr_id_ref=style, inherit_style=False)
            if page_break and index == 0: para.element.set('pageBreak', '1')
        preview.append(text)

    for block in source.iter_inner_content():
        if isinstance(block, Paragraph):
            style_name = block.style.name
            style = title_style if style_name == 'Title' else heading_style if style_name == 'Heading 1' else subheading_style if style_name.startswith('Heading') else body_style
            text = ('• ' if style_name == 'List Bullet' else '') + block.text
            paragraph(text, style, bool(block.paragraph_format.page_break_before))
        elif isinstance(block, Table):
            count = len(block.columns)
            if count == 0 or not block.rows: continue
            cell_width = max(1, width // count)
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            heights = [max(3000, max(_lines(text, cell_width - 800, 9) for text in row) * 1350 + 800) for row in rows]
            table = target.add_table(len(rows), count, width=width, height=sum(heights), char_pr_id_ref=table_style)
            table.element.set('repeatHeader', '1')
            for row_index, row in enumerate(rows):
                for column, text in enumerate(row):
                    cell = table.cell(row_index, column)
                    cell.set_size(width=cell_width, height=heights[row_index])
                    if row_index == 0: cell.element.set('header', '1')
                    table.set_cell_text(row_index, column, text, split_paragraphs=True)
                preview.append(' | '.join(row))
    target.package.set_part('Preview/PrvText.txt', '\n'.join(preview).encode('utf-8'))
    output = BytesIO(); target.save_to_stream(output)
    data = output.getvalue()
    # Verify the archive, XML schemas and reopening with the HWPX reader, not just the filename.
    report = validate_editor_open_safety(data)
    if report.blocking_package_errors or not report.reopen_ok or report.document_validation_error or report.validate_document.issues:
        raise ValueError('HWPX 문서 구조 검증에 실패했습니다.')
    return data
