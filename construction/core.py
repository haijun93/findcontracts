from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED

from .forms import FORM_LABELS, EXTRA_DOCUMENTS, validate_forms, form_warnings, sections as form_sections, survey_table

DOCUMENTS = [('specification', '시방서'), ('plan', '추진계획서')] + EXTRA_DOCUMENTS
DOCUMENT_LABELS = dict(DOCUMENTS)

MAX_FILE = 20 * 1024 * 1024
MAX_TEXT = 200_000
MISSING = '[확인 필요]'


class DocumentError(ValueError):
    pass


def invalid_characters(value):
    return any((ord(c) < 32 and c not in '\n\r\t') or 0xD800 <= ord(c) <= 0xDFFF or ord(c) in (0xFFFE, 0xFFFF) for c in value)


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def extract_document(data, filename, role):
    """Read document contents as data; never execute document instructions or macros."""
    if not data or len(data) > MAX_FILE:
        raise DocumentError('파일은 비어 있지 않은 20MB 이하 파일이어야 합니다.')
    suffix = Path(filename).suffix.lower()
    filename = filename.replace('\\', '/').split('/')[-1][:180]
    pages = []
    try:
        if suffix == '.pdf':
            import pymupdf
            with pymupdf.open(stream=data, filetype='pdf') as document:
                if document.needs_pass or not 1 <= len(document) <= 50:
                    raise DocumentError('암호가 없는 1~50쪽 PDF를 사용하세요.')
                for index, page in enumerate(document):
                    pages.append({'page': index + 1, 'text': page.get_text(sort=True)})
        elif suffix == '.txt':
            try:
                text = data.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = data.decode('cp949')
            pages = [{'page': 1, 'text': text}]
        elif suffix in ('.docx', '.xlsx'):
            with ZipFile(BytesIO(data)) as archive:
                if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                    raise DocumentError('압축 해제 크기가 너무 큰 문서입니다.')
            if suffix == '.docx':
                from docx import Document
                from docx.table import Table
                from docx.text.paragraph import Paragraph
                doc = Document(BytesIO(data))
                lines = []
                for block in doc.iter_inner_content():
                    if isinstance(block, Paragraph):
                        lines.append(block.text)
                    elif isinstance(block, Table):
                        lines.extend('  '.join(cell.text for cell in row.cells) for row in block.rows)
                pages = [{'page': 1, 'text': '\n'.join(lines)}]
            else:
                from openpyxl import load_workbook
                book = load_workbook(BytesIO(data), read_only=True, data_only=True)
                try:
                    for index, sheet in enumerate(book.worksheets):
                        if index >= 30 or sheet.max_row > 5000 or sheet.max_column > 100:
                            raise DocumentError('Excel은 30개 시트, 시트당 5천 행·100열 이하로 입력하세요.')
                        lines = ['  '.join(str(v) if v is not None else '' for v in row) for row in sheet.iter_rows(values_only=True)]
                        pages.append({'page': index + 1, 'text': '\n'.join(lines), 'sheet': sheet.title})
                finally:
                    book.close()
        else:
            raise DocumentError('PDF, DOCX, XLSX, TXT 파일만 지원합니다.')
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError('문서를 읽지 못했습니다. 파일 형식·암호·손상 여부를 확인하세요.') from error
    if sum(len(p['text']) for p in pages) > MAX_TEXT:
        raise DocumentError('문서의 텍스트가 너무 깁니다. 공사 관련 부분만 입력하세요.')
    if not any(len(clean(p['text'])) >= 10 for p in pages):
        raise DocumentError('텍스트가 없는 스캔 문서입니다. OCR 처리한 파일이나 TXT 문서를 입력하세요.')
    blank = [p['page'] for p in pages if len(clean(p['text'])) < 10]
    return {'role': role, 'filename': filename, 'sha256': sha256(data).hexdigest(),
            'pages': pages, 'warnings': [f'{filename}: {blank}쪽은 텍스트를 읽지 못했습니다. 그림·사진은 원본을 확인하세요.'] if blank else []}


def analyze(quote, request_doc):
    facts = {key: '' for key in ('title', 'purpose', 'location', 'scope', 'duration', 'department',
                                'budget_code', 'contract_method', 'schedule')}
    facts.update({key: '' for key in FORM_LABELS})
    facts.update(comparisons=[])
    facts.update(items=[], sites=[], budget={key: None for key in ('subtotal', 'adjustment', 'vat', 'total')})
    evidence = []

    def record(field, value, doc, page, excerpt):
        if value not in ('', None):
            evidence.append({'field': field, 'value': value, 'source': f'{doc["filename"]} · {page}쪽/시트', 'excerpt': clean(excerpt)[:800]})

    for page in quote['pages']:
        text = page['text']
        m = re.search(r'견\s*적\s*서\s*[(（]([^\n)）]+)[)）]', text)
        if m and not facts['title']:
            facts['title'] = clean(m[1]); record('title', facts['title'], quote, page['page'], m[0])
        if not facts['business_number']:
            numbers = set(re.findall(r'(?<!\d)\d{3}-\d{2}-\d{5}(?!\d)', text))
            if len(numbers) == 1:
                facts['business_number'] = next(iter(numbers))
                record('business_number', facts['business_number'], quote, page['page'], facts['business_number'])
        for raw in text.splitlines():
            party = re.search(r'상호\s+(.+?)\s+대표\s+(.+?)(?=\s{2,}|$)', raw)
            if party and not facts['supplier']:
                facts['supplier'] = clean(party[1])
                facts['representative'] = clean(re.sub(r'\((?:인|서명|날인)\)', '', party[2]))
                record('supplier', facts['supplier'], quote, page['page'], raw)
                record('representative', facts['representative'], quote, page['page'], raw)
            line = raw.strip()
            amount = re.findall(r'(?<![\d.,])\d{1,3}(?:,\d{3})+(?![\d.,])', line)
            label = re.sub(r'\s+', '', line)
            key = None
            if '합계' in label:
                key = 'total'
            elif '소계' in label:
                key = 'subtotal'
            elif '절사' in label or '할인' in label:
                key = 'adjustment'
            elif '부가가치세' in label or '부가세' in label:
                key = 'vat'
            if key and amount:
                # Header total may be followed by a date; use the explicit currency number.
                value = int(amount[-1].replace(',', ''))
                if key == 'adjustment': value = -value
                previous = facts['budget'][key]
                if previous is not None and previous != value:
                    record('budget_conflict.' + key, value, quote, page['page'], line)
                else:
                    facts['budget'][key] = value
                    record('budget.' + key, value, quote, page['page'], line)
                continue
            m = re.match(r'^\s*\d{1,3}[.)]?\s+(.+)', line)
            if not m or not re.search(r'[가-힣A-Za-z]', m[1]):
                continue
            parts = re.split(r'\s{2,}|\t+', m[1].strip())
            if not parts or parts[0] in ('품목', '구분'):
                continue
            item = {'name': clean(parts[0]), 'specification': '', 'unit': '', 'quantity': '',
                    'unit_price': None, 'amount': None,
                    'source': f'{quote["filename"]} · {page["page"]}쪽/시트'}
            # Only map unambiguous rows: name, unit, quantity, unit price, amount.
            if len(parts) == 5 and re.fullmatch(r'[가-힣A-Za-z²³/]+', parts[1]) and re.fullmatch(r'\d+(?:\.\d+)?', parts[2]):
                item['unit'], item['quantity'] = parts[1:3]
                if all(re.fullmatch(r'\d[\d,]*', value) for value in parts[3:5]):
                    item['unit_price'], item['amount'] = [int(v.replace(',', '')) for v in parts[3:5]]
            elif len(parts) > 1:
                item['specification'] = ' / '.join(parts[1:])
            facts['items'].append(item)
            record('items', item['name'], quote, page['page'], line)

    for doc in (quote, request_doc):
        for page in doc['pages']:
            for raw in page['text'].splitlines():
                for label, field in [('공사명', 'title'), ('추진목적', 'purpose'), ('예산과목', 'budget_code')]:
                    match = re.match(r'^\s*[□ㅇ○-]?\s*' + label + r'\s*[:：]\s*(.+)$', raw)
                    if match and not facts[field]:
                        facts[field] = clean(match[1])
                        record(field, facts[field], doc, page['page'], raw)

    for page in request_doc['pages']:
        text = page['text']
        for raw in text.splitlines():
            line = clean(raw)
            if not facts['title'] and re.search(r'(?:의뢰|요청)\s*$', line):
                facts['title'] = re.sub(r'\s*(?:의뢰|요청)\s*$', '', line)
                record('title', facts['title'], request_doc, page['page'], line)
            m = re.search(r'(?:과업내용|공사내용|작업내용)\s*[:：]\s*(.+)', line)
            if m:
                facts['scope'] += ('\n' if facts['scope'] else '') + m[1]
                record('scope', m[1], request_doc, page['page'], line)
            m = re.search(r'(?:위치도|공사위치|공사장소|장소)\s*[:：]\s*(.+)', line)
            if m and not facts['location']:
                facts['location'] = m[1]; record('location', m[1], request_doc, page['page'], line)
            m = re.search(r'(?:공사기간|공기|기간)\s*[:：]\s*(.+)', line)
            if m:
                facts['duration'] = m[1]; record('duration', m[1], request_doc, page['page'], line)
            m = re.search(r'([가-힣]+구\s+[가-힣]+과\s+[가-힣]+팀)', line)
            if m and not facts['department']:
                facts['department'] = m[1]; record('department', m[1], request_doc, page['page'], line)
            # Site lists in the supplied request: number / street address / power source note.
            parts = re.split(r'\s{2,}|\t+', raw.strip())
            if len(parts) >= 2 and re.fullmatch(r'\d{1,3}', parts[0]) and re.search(r'(?:로|길|동)\s*\d', parts[1]):
                site = {'address': clean(parts[1]), 'note': clean(' / '.join(parts[2:])),
                        'source': f'{request_doc["filename"]} · {page["page"]}쪽/시트'}
                facts['sites'].append(site); record('sites', site['address'], request_doc, page['page'], raw)
    if not facts['title']:
        facts['title'] = ''
    conflicts = [f'견적서 내 {e["field"]} 금액이 서로 다릅니다. 원본을 확인하세요.' for e in evidence if e['field'].startswith('budget_conflict.')]
    return {'facts': facts, 'evidence': evidence, 'sources': [quote, request_doc],
            'warnings': quote['warnings'] + request_doc['warnings'] + conflicts + assess(facts)}


def money(value):
    return f'{value:,}원' if value is not None else MISSING


def assess(facts):
    warnings = []
    for field, label in [('title', '공사명'), ('scope', '공사내용'), ('location', '위치'),
                         ('duration', '공사기간'), ('purpose', '추진목적'), ('department', '담당 부서'), ('budget_code', '예산과목'),
                         ('contract_method', '계약방법·근거'), ('schedule', '추진일정')]:
        if not facts.get(field):
            warnings.append(f'{label}: 입력 근거가 없거나 미입력입니다. 문서에 확인 필요로 표시합니다.')
    budget = facts.get('budget', {})
    total = budget.get('total')
    if total is None:
        warnings.append('총액이 확인되지 않았습니다.')
    keys = ('subtotal', 'adjustment', 'vat')
    if total is not None and all(budget.get(k) is not None for k in keys):
        calculated = sum(budget[k] for k in keys)
        if calculated != total:
            warnings.append(f'예산 합계 불일치: 소계+조정액+부가세={money(calculated)}, 총액={money(total)}. 확인하세요.')
    if not facts.get('items'):
        warnings.append('공사내역을 자동 추출하지 못했습니다. 원문을 보고 내역을 입력하세요.')
    for item in facts.get('items', []):
        if item.get('amount') is None:
            warnings.append(f'견적 포함 여부 확인: {item["name"]}의 개별 금액이 없습니다.')
        if item.get('quantity') and item.get('unit_price') is not None and item.get('amount') is not None:
            if Decimal(item['quantity']) * item['unit_price'] != item['amount']:
                warnings.append(f'내역 산술 불일치: {item["name"]}의 수량×단가와 금액이 다릅니다.')
    scope = facts.get('scope', '')
    quantity = re.search(r'(\d+)\s*(?:개소|개)', scope)
    if quantity and facts.get('sites') and int(quantity[1]) != len(facts['sites']):
        warnings.append('의뢰 수량과 위치 목록 개수가 다릅니다.')
    if quantity and len(facts.get('items', [])):
        first = facts['items'][0].get('quantity')
        if first and Decimal(first) != int(quantity[1]):
            warnings.append('의뢰 수량과 견적 첫 항목 수량이 다릅니다.')
    warnings.append('현장사진·도면은 자동 해석하지 않습니다. 원본과 대조하고 현장 조건을 확인하세요.')
    return warnings


def validate_facts(value):
    if not isinstance(value, dict):
        raise DocumentError('검토 내용을 입력하세요.')
    result = {}
    for key in ('title', 'purpose', 'location', 'scope', 'duration', 'department', 'budget_code', 'contract_method', 'schedule'):
        text = value.get(key, '')
        if not isinstance(text, str) or len(text) > 5000 or invalid_characters(text):
            raise DocumentError('문자 필드는 5천 자 이하로 입력하세요.')
        result[key] = text.strip()
    if not result['title'] or not result['scope']:
        raise DocumentError('공사명과 공사내용을 확인하여 입력하세요.')

    def number(v, signed=False):
        if v in ('', None): return None
        if isinstance(v, bool) or not re.fullmatch(r'-?(?:\d+|\d{1,3}(?:,\d{3})+)', str(v)):
            raise DocumentError('금액은 원 단위 정수로 입력하세요.')
        n = int(str(v).replace(',', ''))
        if abs(n) > 10**15 or (not signed and n < 0):
            raise DocumentError('금액 범위를 확인하세요.')
        return n
    budget = value.get('budget', {})
    if not isinstance(budget, dict): raise DocumentError('예산 입력 형식을 확인하세요.')
    result['budget'] = {k: number(budget.get(k), signed=k == 'adjustment') for k in ('subtotal', 'adjustment', 'vat', 'total')}
    for key, fields in [('items', ('name','specification','unit','quantity','source')), ('sites', ('address','note','source'))]:
        rows = value.get(key, [])
        if not isinstance(rows, list) or len(rows) > 200:
            raise DocumentError('내역과 위치는 각각 200개 이하로 입력하세요.')
        result[key] = []
        for row in rows:
            if not isinstance(row, dict): raise DocumentError('표의 입력 형식을 확인하세요.')
            target = {}
            for field in fields:
                v = row.get(field, '')
                if not isinstance(v, str) or len(v) > 1000 or invalid_characters(v):
                    raise DocumentError('표의 문자는 1천 자 이하로 입력하세요.')
                target[field] = v.strip()
            if key == 'items':
                q = target['quantity']
                if q and (not re.fullmatch(r'\d{1,9}(?:\.\d{1,6})?', q) or Decimal(q) <= 0):
                    raise DocumentError('수량은 양수로 입력하세요. 미확인은 비워 두세요.')
                target['unit_price'] = number(row.get('unit_price'))
                target['amount'] = number(row.get('amount'))
            result[key].append(target)
    result.update(validate_forms(value, invalid_characters, DocumentError, number))
    return result


def draft_sections(facts, kind):
    if kind not in ('specification', 'plan'):
        return form_sections(facts, kind, MISSING, money)
    f = lambda key: facts.get(key) or MISSING
    scope = f('scope')
    general = [
        ('적용범위', f'본 문서는 {f("title")}의 검토용 초안이다. 적용범위는 다음과 같으며 계약도서와 현장 확인 후 확정한다.\n{scope}'),
        ('공사기간 및 현장관리', f'공사기간: {f("duration")}\n착수 전 작업 순서·장비·보행 동선·시설물 보호계획을 담당 부서와 협의한다. 인접 시설 및 매설물은 현장 조사 후 보호한다.'),
        ('안전·환경 관리', '작업구역을 구획하고 보행자·차량의 안전한 통행을 확보한다. 소음·분진을 저감하고 폐기물은 분리 수거하여 적정 처리한다. 위험이 발견되면 작업을 중지하고 감독자에게 보고한다.'),
    ]
    special = []
    # These are proposed execution checks, not extracted facts or fixed engineering specifications.
    if re.search(r'전기|VCTF|후렉|분전|전선', scope + ' '.join(i['name'] for i in facts['items']), re.I):
        special.append(('전기 관련 작업 검토안', '전기 작업 포함 시 전원 공급 경로와 공용 회로를 확인하고 시설 관리자와 정전 범위를 협의한다. 적격 작업자가 차단·무전압 확인 후 작업하며, 존치 회로를 보호한다. 철거 후 잔여 배선·관로의 절연·방수·마감 상태와 존치 시설의 정상 작동을 확인한다. 회로 구성과 시험 방법은 현장 확인 후 확정한다.'))
    if re.search(r'철거|해체', scope):
        special.append(('철거 작업 검토안', '철거 대상과 존치 대상을 현장에서 구분하고 전도·낙하 방지 조치를 수립한다. 지주·기초의 철거 범위와 운반·반출 방법은 사전 협의한다. 철거 후 개구부와 단차를 정리하고 복구 범위를 확인한다.'))
    if re.search(r'이설|이전|설치', scope):
        special.append(('이전·설치 작업 검토안', '이전·설치 위치와 구조·기초·고정 방법은 현장 및 도면 확인 후 결정한다. 반입·운반 중 시설을 보호하고 설치 후 수평·수직·고정 상태와 기능을 확인한다. 규격·이격거리·기초치수는 입력 도서에 없는 경우 임의로 정하지 않는다.'))
    if kind == 'specification':
        return [('1. 일반시방서', general), ('2. 특별시방서', special or [('공종별 사항', '입력된 공사내역별 자재·공법·규격·허용오차 및 검사기준은 설계도서와 현장 확인 후 확정한다.')]),
                ('3. 검사 및 제출자료', [('준공 확인', '작업 전·중·후 사진, 작업내역, 폐기물 처리 증빙 및 해당 시험·검사 결과를 제출한다. 내역·위치별 완료 상태와 현장 정리를 확인하고 미비 사항을 보완한다.')])]
    return [('Ⅰ. 추진개요', [('목적', f('purpose')), ('추진근거·의뢰', '입력한 의뢰서 및 견적서에 따른 검토 초안이다. 추진근거 문서번호 및 법적 요건: ' + MISSING)]),
            ('Ⅱ. 추진계획', [('공사개요', f'공사명: {f("title")}\n위치: {f("location")}\n담당: {f("department")}\n기간: {f("duration")}\n공사내용: {scope}'), ('계약방법', f('contract_method'))]),
            ('Ⅲ. 소요예산', [('예산과목', f('budget_code'))]),
            ('Ⅳ. 추진일정', [('일정', f('schedule'))]),
            ('Ⅴ. 행정사항', [('확인·협의사항', '담당 부서·시설 관리자와 공사 범위 및 통행·정전 계획을 협의한다. 계약방법, 계약심사·감사 대상 여부, 인허가·신고 및 적용 법령은 담당자가 검토하여 확정한다.\n붙임: 시방서(초안) 및 입력 견적서·의뢰서 참조.')])]


def render_docx(facts, kind, warnings):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = '맑은 고딕'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    doc.styles['Normal'].font.size = Pt(10)
    label = DOCUMENT_LABELS[kind]
    doc.add_heading(f'{facts["title"]}\n{label} (초안)', 0)
    doc.add_paragraph('담당자 검토 전 초안 · 미입력 항목은 [확인 필요] · 일반 시공 문구는 제안사항이며 현장 적합성 확인 필요')
    for heading, entries in draft_sections(facts, kind):
        doc.add_heading(heading, 1)
        for title, content in entries:
            doc.add_heading(title, 2)
            doc.add_paragraph(content)
        if (kind == 'specification' and heading.startswith('2.')) or (kind == 'plan' and heading.startswith('Ⅱ.')):
            add_tables(doc, facts, budget=False)
        if kind == 'plan' and heading.startswith('Ⅲ.'):
            add_tables(doc, facts, budget=True)
        if kind == 'price_survey' and heading.startswith('2.'):
            doc.add_heading('비교견적 조사내역 (단위: 원)', 1)
            headers, rows = survey_table(facts, MISSING, money)
            table = doc.add_table(rows=1, cols=len(headers)); table.style = 'Table Grid'
            for cell, text in zip(table.rows[0].cells, headers): cell.text = text
            for row in rows:
                for cell, text in zip(table.add_row().cells, row): cell.text = text
    doc.add_heading('검토 필요사항', 1)
    for warning in warnings:
        doc.add_paragraph(warning, style='List Bullet')
    doc.add_paragraph('첨부 사진·도면은 원본 참조. 예시 문서의 금액·공기·법령판단·결재정보는 이 초안에 전용하지 않았다.')
    footer = section.footer.paragraphs[0]
    footer.add_run('검토용 초안 | ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    output = BytesIO(); doc.save(output)
    return output.getvalue()


def add_tables(doc, facts, budget=False):
    if budget:
        table = doc.add_table(rows=1, cols=2); table.style = 'Table Grid'
        table.rows[0].cells[0].text = '예산 구분'; table.rows[0].cells[1].text = '금액(원)'
        for key, label in [('subtotal','소계'),('adjustment','조정액(절사·할인 등)'),('vat','부가가치세'),('total','총액')]:
            cells = table.add_row().cells; cells[0].text = label; cells[1].text = money(facts['budget'][key])
        return
    doc.add_heading('공사내역 (입력 및 검토값)', 2)
    table = doc.add_table(rows=1, cols=6); table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['품목','규격','단위','수량','단가(원)','금액(원)']): cell.text = text
    for item in facts['items']:
        cells = table.add_row().cells
        values = [item['name'], item['specification'] or MISSING, item['unit'] or MISSING,
                  item['quantity'] or MISSING, money(item['unit_price']), money(item['amount'])]
        for cell, text in zip(cells, values): cell.text = text
    if not facts['items']: doc.add_paragraph('세부내역: ' + MISSING)
    doc.add_heading('위치별 작업대상', 2)
    table = doc.add_table(rows=1, cols=3); table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['번호','위치','현장 메모']): cell.text = text
    for index, site in enumerate(facts['sites'], 1):
        for cell, text in zip(table.add_row().cells, [str(index), site['address'], site['note'] or MISSING]): cell.text = text
    if not facts['sites']: doc.add_paragraph('위치별 목록: ' + MISSING)


def render_markdown(facts, kind, warnings):
    text = [f'# {facts["title"]} — {DOCUMENT_LABELS[kind]} (초안)',
            '담당자 검토 전 초안. 미입력 항목은 [확인 필요]. 일반 시공 문구는 검토 제안사항입니다.']
    for heading, entries in draft_sections(facts, kind):
        text.append('## ' + heading)
        for title, content in entries: text.extend(['### ' + title, content])
    if kind in ('specification', 'plan'):
        text.append('## 공사내역')
        for i in facts['items']:
            text.append(f'- {i["name"]} / {i["specification"] or MISSING} / {i["quantity"] or MISSING}{i["unit"]} / 단가 {money(i["unit_price"])} / 금액 {money(i["amount"])}')
        text.append('## 위치 목록')
        text.extend(f'- {i["address"]} / {i["note"] or MISSING}' for i in facts['sites'])
        if kind == 'plan':
            text.append('## 예산 상세')
            text.extend(f'- {label}: {money(facts["budget"][key])}' for key,label in [('subtotal','소계'),('adjustment','조정액'),('vat','부가세'),('total','총액')])
    if kind == 'price_survey':
        headers, rows = survey_table(facts, MISSING, money)
        text.append('## 비교견적 조사내역')
        text.append(' / '.join(headers))
        text.extend('- ' + ' / '.join(row) for row in rows)
    text.extend(['## 검토 필요사항', *['- ' + w for w in warnings]])
    return '\n\n'.join(text)


def generate_package(facts, provenance=None):
    facts = validate_facts(facts)
    warnings = assess(facts)
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        for kind, title in DOCUMENTS:
            label = title.replace(' ', '_') + '_초안'
            specific_warnings = warnings + form_warnings(facts, kind)
            archive.writestr(label + '.docx', render_docx(facts, kind, specific_warnings))
            archive.writestr(label + '.md', render_markdown(facts, kind, specific_warnings))
        manifest = {'status': '사용자 검토값 기반 초안; 자동 검증·승인된 계약문서가 아님',
                    'facts': facts, 'warnings': warnings,
                    'form_warnings': {kind: form_warnings(facts, kind) for kind, _ in EXTRA_DOCUMENTS},
                    'source_extraction': provenance or {},
                    'note': 'source_extraction은 브라우저가 전달한 추출 기록이며 전자서명·원본 인증이 아닙니다.'}
        try:
            record = json.dumps(manifest, ensure_ascii=False, indent=2).encode('utf-8')
        except (TypeError, UnicodeError) as error:
            raise DocumentError('추출 근거에 저장할 수 없는 문자가 포함되어 있습니다.') from error
        archive.writestr('추출근거_검토기록.json', record)
    output.seek(0)
    return output
